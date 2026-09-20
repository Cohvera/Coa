"""Read customer XLSX as data. Never execute formulas or embedded instructions."""
from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
from decimal import Decimal, InvalidOperation
import openpyxl
from openpyxl.utils import column_index_from_string

MAX_BYTES = 10 * 1024 * 1024


def number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, bool):
        raise ValueError("Een getal is vereist.")
    text = str(value).strip().replace("\u00a0", "").replace(" ", "")
    if "," in text:
        text = text.replace(".", "").replace(",", ".")
    try:
        result = Decimal(text)
    except InvalidOperation:
        raise ValueError("Ongeldig getal.")
    if not result.is_finite() or result < 0 or result > 100000000:
        raise ValueError("Getal moet tussen 0 en 100.000.000 liggen.")
    return str(result)


def unit(value):
    result = str(value or "").strip().lower().rstrip(".")
    return {"stuk": "st", "stuks": "st", "pcs": "st", "lm": "m"}.get(result, result)


def normalize(value):
    return re.sub(r"\s+", " ", str(value or "").strip().casefold())


def load(data, cached=True):
    if len(data) > MAX_BYTES:
        raise ValueError("Bestand is groter dan 10 MB.")
    try:
        with ZipFile(BytesIO(data)) as archive:
            if len(archive.infolist()) > 2000 or sum(x.file_size for x in archive.infolist()) > 50 * 1024 * 1024:
                raise ValueError("Uitgepakt bestand is te groot.")
        return openpyxl.load_workbook(BytesIO(data), read_only=True, data_only=cached, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("Dit bestand is geen leesbare XLSX.") from exc


def inspect(data):
    workbook = load(data)
    try:
        sheets = []
        for sheet in workbook:
            if sheet.max_row > 20000 or sheet.max_column > 100:
                raise ValueError("Maximaal 20.000 rijen en 100 kolommen per werkblad.")
            sample = [[str(v)[:300] if v is not None else "" for v in row]
                      for row in sheet.iter_rows(min_row=1, max_row=min(sheet.max_row, 45), values_only=True)]
            structured_layout = any(len(r) > 11 and r[9] == "#" and r[10] == "EH" for r in sample)
            mapping = {"description": "D", "quantity": "J", "unit": "K", "code": "A", "context": "B", "start": 35} if structured_layout and "ELEKTR" in sheet.title.upper() else {"description": "A", "quantity": "B", "unit": "C", "code": "", "context": "", "start": 2}
            sheets.append({"name": sheet.title, "rows": sheet.max_row, "sample": sample, "mapping": mapping})
        return sheets
    finally:
        workbook.close()


def parse(data, sheet_name, mapping):
    sheets = inspect(data)
    if sheet_name not in [s["name"] for s in sheets]:
        raise ValueError("Werkblad bestaat niet.")
    indexes = {}
    for field in ("description", "quantity", "unit", "code", "context"):
        col = str(mapping.get(field, "")).strip().upper()
        if not col and field in ("code", "context"):
            indexes[field] = None
            continue
        try:
            index = column_index_from_string(col) - 1
        except ValueError:
            raise ValueError(f"Ongeldige kolom voor {field}.")
        if index >= 100:
            raise ValueError("Kolom buiten bereik.")
        indexes[field] = index
    if len(set(i for i in indexes.values() if i is not None)) != len([i for i in indexes.values() if i is not None]):
        raise ValueError("Kies verschillende kolommen.")
    start = int(mapping.get("start", 2))
    if start < 1 or start > 20000:
        raise ValueError("Ongeldige eerste rij.")
    cached, formulas = load(data), load(data, False)
    result, notes = [], []
    section = group = context = ""
    try:
        for row_number, (row, formula_row) in enumerate(zip(cached[sheet_name].values, formulas[sheet_name].values), 1):
            def get(field):
                idx = indexes[field]
                return row[idx] if idx is not None and idx < len(row) else None
            description, code, ctx = get("description"), get("code"), get("context")
            # Section codes identify hierarchy; quantities are already totals in the source.
            if code and re.fullmatch(r"\d{4}\.\d{2}", str(code)):
                section, group, context = str(ctx or ""), "", ""
            elif code and re.fullmatch(r"\d{4}\.\d{2}\.\d{2}", str(code)):
                group, context = str(ctx or ""), ""
            elif ctx:
                context = str(ctx).strip()
            if row_number < start:
                if ctx and not code:
                    notes.append({"row": row_number, "text": str(ctx)})
                continue
            if description is None or str(description).strip() == "":
                if ctx and not code:
                    notes.append({"row": row_number, "text": str(ctx)})
                continue
            if normalize(description) in ("omschrijving", "description", "artikel", "totaal", "subtotal", "subtotaal"):
                continue
            issue = ""
            try:
                quantity = number(get("quantity"))
            except ValueError:
                quantity, issue = None, "Ongeldige hoeveelheid; controleer de bron."
            qi = indexes["quantity"]
            if quantity is None and qi < len(formula_row) and str(formula_row[qi]).startswith("="):
                issue = "Formule zonder opgeslagen uitkomst; vul de hoeveelheid in."
            if quantity is None and not issue:
                issue = "Hoeveelheid ontbreekt."
            desc = str(description).strip()
            match_text = f"{context} / {desc}" if normalize(desc) in ("gemene delen", "privatieven") and context else desc
            result.append({"id": row_number, "source_row": row_number, "sheet": sheet_name,
                           "section": section, "group": group, "context": context, "description": desc,
                           "match_text": match_text, "quantity": quantity, "source_quantity": quantity,
                           "unit": unit(get("unit")), "source_unit": unit(get("unit")), "catalog_id": None,
                           "excluded": False, "reason": "", "issue": issue})
        if not result:
            raise ValueError("Geen detailregels gevonden. Controleer werkblad en kolommen.")
        return {"lines": result, "notes": notes}
    finally:
        cached.close()
        formulas.close()
